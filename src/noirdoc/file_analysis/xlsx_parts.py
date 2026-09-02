"""Pseudonymize and reveal XLSX parts that live outside the cell grid.

``pseudonymize_xlsx_smart`` only rewrites data cells. A plain openpyxl
load→save then preserves everything else verbatim, so a redacted workbook
still leaked: ``docProps/core.xml`` (creator, lastModifiedBy, title, …),
``docProps/custom.xml``, cell comments (author + text), sheet headers/footers,
pivot caches — whose records are a literal snapshot of the source rows, so
Excel keeps displaying the original names in the pivot even after the sheet
cells were pseudonymized — plus chart string caches and titles (Excel always
writes a cache of the charted cells), hyperlink tooltips / display strings and
``mailto:`` targets, AutoFilter criteria, conditional-formatting text
literals, data-validation inline lists and prompts, and pivot-table captions
and label filters.

Design: **one** slot enumerator (:func:`iter_text_slots`) feeds **both**
directions — :func:`pseudonymize_workbook_parts` and
:func:`reidentify_workbook_parts` — so the redact and reveal walkers cannot
drift apart. Every value is pseudonymized through the shared mapper (never
blanked), which keeps ``noirdoc reveal`` a full round-trip.

Surface names reported in ``classifications`` never carry file-provided
strings other than column/field/property *names*: sheets are referred to by
index, threaded-comment persons by ordinal. Classifications are logged.
"""

from __future__ import annotations

import asyncio
import io
import re
import zipfile
from collections.abc import Iterator, Mapping
from dataclasses import dataclass, field
from typing import TYPE_CHECKING, Any, Literal
from urllib.parse import quote, unquote

import structlog

from noirdoc.file_analysis.xlsx_inference import classify_by_sample, infer_entity_type

if TYPE_CHECKING:
    from openpyxl.workbook.workbook import Workbook

    from noirdoc.detection.base import DetectedEntity
    from noirdoc.file_analysis.xlsx_inference import DetectorLike
    from noirdoc.pseudonymization.mapper import PseudonymMapper
    from noirdoc.reidentification.engine import ReidentificationEngine

logger = structlog.get_logger()

_PLACEHOLDER = re.compile(r"<<[A-Z_]+_\d+>>")
_PLACEHOLDER_TYPE = re.compile(r"<<([A-Z_]+)_\d+>>")
# Excel writes a legacy mirror of every threaded comment with this synthetic author.
_THREADED_MIRROR_AUTHOR = re.compile(r"^tc=\{[0-9A-Fa-f-]{36}\}$")

# core.xml fields that are a person (or account) by definition — never rely on NER for them.
_CORE_AUTHOR_FIELDS = ("creator", "lastModifiedBy")
# core.xml free-text fields — run through the detector like body text.
_CORE_FREE_FIELDS = (
    "title",
    "subject",
    "description",
    "keywords",
    "category",
    "contentStatus",
    "identifier",
)
# (attribute on ws.HeaderFooter, surface label)
_HEADER_FOOTER_ITEMS = (
    ("oddHeader", "header.odd"),
    ("oddFooter", "footer.odd"),
    ("evenHeader", "header.even"),
    ("evenFooter", "footer.even"),
    ("firstHeader", "header.first"),
    ("firstFooter", "footer.first"),
)
_HEADER_FOOTER_PARTS = ("left", "center", "right")

# ``author``: a person/account by definition — forced mapping, never NER.
# ``free``:   free text — offset-based substitution of detected entities.
# ``pivot``:  a cached copy of classified column data — replaced as a whole
#             when its field is classified or the value is already mapped.
# ``filter``: a criterion/caption mirroring cell values — replaced as a whole
#             when the mapper already knows it (the cell pass ran first),
#             otherwise treated like free text so wildcards survive.
SlotKind = Literal["author", "free", "pivot", "filter"]
FieldKey = tuple[int, int]  # (pivot cache ordinal, cacheField index)

# An Excel formula that is exactly one double-quoted string literal.
_QUOTED_LITERAL = re.compile(r'^"(?:[^"]|"")*"$')


def _formula_quote(value: str) -> str:
    return '"' + value.replace('"', '""') + '"'


def _formula_unquote(formula: str) -> str:
    return formula[1:-1].replace('""', '"')


@dataclass
class _Slot:
    """One string-bearing attribute somewhere in the workbook object model."""

    surface: str  # human-readable location, e.g. "core.creator"
    kind: SlotKind
    obj: Any  # openpyxl object holding the string
    attr: str  # attribute name on ``obj``
    field_key: FieldKey | None = None  # pivot items: which cache field the item belongs to
    field_name: str | None = None  # pivot items: the cacheField name (= source column header)
    count: bool = True  # False for a duplicate object of an already-counted pivot cache

    @property
    def value(self) -> str | None:
        """Current string, or ``None`` when the slot is empty / not text."""
        raw = getattr(self.obj, self.attr, None)
        if isinstance(raw, str) and raw.strip():
            return raw
        return None

    def set(self, value: str) -> None:
        setattr(self.obj, self.attr, value)


class _HeaderFooterRaw:
    """A header/footer part exposed as its raw code string.

    openpyxl's ``&"font"`` parse is greedy up to the LAST double quote, so text
    between two font codes (a bold name, say) lands in ``.font`` and ``.text``
    never sees it. Working on the raw string reaches it; writing back re-parses
    the same way, which is what the writer serializes from.
    """

    def __init__(self, part: Any) -> None:
        self._part = part

    @property
    def raw(self) -> str:
        return str(self._part)

    @raw.setter
    def raw(self, value: str) -> None:
        from openpyxl.worksheet.header_footer import _HeaderFooterPart

        parsed = _HeaderFooterPart.from_str(value)
        for attr in ("text", "font", "size", "color"):
            setattr(self._part, attr, getattr(parsed, attr))


class _MailtoTarget:
    """A hyperlink ``mailto:`` target exposed as its percent-decoded address.

    ``<`` and ``>`` are illegal URI characters, so a raw ``<<EMAIL_1>>`` in the
    relationship ``Target`` would risk Excel's repair prompt. The placeholder is
    therefore percent-encoded on write (``mailto:%3C%3CEMAIL_1%3E%3E`` — RFC
    6068 explicitly allows percent-encoding in mailto URIs) and decoded again by
    the getter, so redact and reveal both see the plain placeholder.
    """

    def __init__(self, link: Any) -> None:
        self._link = link

    @property
    def address(self) -> str | None:
        target = getattr(self._link, "target", None)
        if isinstance(target, str) and target.lower().startswith("mailto:"):
            return unquote(target[len("mailto:") :])
        return None

    @address.setter
    def address(self, value: str) -> None:
        # "@?=&" keeps the address and any ?subject= query structure intact.
        self._link.target = "mailto:" + quote(value, safe="@?=&")


class _ListItem:
    """One entry of a plain string list (e.g. ``Filters.filter``)."""

    def __init__(self, seq: list[Any], index: int) -> None:
        self._seq = seq
        self._index = index

    @property
    def item(self) -> Any:
        return self._seq[self._index]

    @item.setter
    def item(self, value: str) -> None:
        self._seq[self._index] = value


class _RuleText:
    """A conditional-formatting rule's ``text`` plus its echo in the formulas.

    Excel stores a containsText criterion twice: as ``text="Anna"`` and inside
    the rule formula (``SEARCH("Anna",A1)``). Writing through this wrapper
    rewrites the quoted literal in every formula in lockstep, in both
    directions — redact syncs the placeholder in, reveal syncs it back out.
    """

    def __init__(self, rule: Any) -> None:
        self._rule = rule

    @property
    def text(self) -> Any:
        return self._rule.text

    @text.setter
    def text(self, value: str) -> None:
        old = self._rule.text
        if isinstance(old, str) and old:
            needle, replacement = _formula_quote(old), _formula_quote(value)
            self._rule.formula = [f.replace(needle, replacement) for f in self._rule.formula]
        self._rule.text = value


class _FormulaLiteral:
    """A rule formula that is exactly one quoted string literal (cellIs equal …)."""

    def __init__(self, rule: Any, index: int) -> None:
        self._rule = rule
        self._index = index

    @property
    def literal(self) -> str | None:
        formula = self._rule.formula[self._index]
        if isinstance(formula, str) and _QUOTED_LITERAL.match(formula):
            return _formula_unquote(formula)
        return None

    @literal.setter
    def literal(self, value: str) -> None:
        formulas = list(self._rule.formula)
        formulas[self._index] = _formula_quote(value)
        self._rule.formula = formulas


class _InlineListItem:
    """One entry of a data validation's inline list (``formula1='"Anna,Ben"'``).

    List entries cannot contain commas in the XLSX format, so splitting and
    re-joining on ``,`` is lossless — and placeholders contain no commas, so
    the entry count is stable between redact and reveal.
    """

    def __init__(self, dv: Any, index: int) -> None:
        self._dv = dv
        self._index = index

    def _items(self) -> list[str] | None:
        formula = self._dv.formula1
        if isinstance(formula, str) and _QUOTED_LITERAL.match(formula):
            return _formula_unquote(formula).split(",")
        return None

    @property
    def item(self) -> str | None:
        items = self._items()
        if items is not None and self._index < len(items):
            return items[self._index]
        return None

    @item.setter
    def item(self, value: str) -> None:
        items = self._items()
        if items is None or self._index >= len(items):
            return
        items[self._index] = value
        self._dv.formula1 = _formula_quote(",".join(items))


@dataclass
class PartsResult:
    """What :func:`pseudonymize_workbook_parts` found (and, unless counting only, replaced)."""

    entity_types: dict[str, int] = field(default_factory=dict)
    classifications: dict[str, str] = field(default_factory=dict)

    @property
    def entity_count(self) -> int:
        return sum(self.entity_types.values())

    def _add(self, entity_type: str, count: int = 1) -> None:
        self.entity_types[entity_type] = self.entity_types.get(entity_type, 0) + count


def _header_footer_slots(sheet: Any, where: str) -> Iterator[_Slot]:
    for hf_attr, label in _HEADER_FOOTER_ITEMS:
        item = getattr(sheet.HeaderFooter, hf_attr)
        for part_name in _HEADER_FOOTER_PARTS:
            part = getattr(item, part_name)
            surface = f"{label}.{part_name}@{where}"
            font = getattr(part, "font", None)
            if isinstance(font, str) and '"' in font:
                yield _Slot(surface, "free", _HeaderFooterRaw(part), "raw")
            else:
                yield _Slot(surface, "free", part, "text")


def _autofilter_slots(
    auto_filter: Any,
    surface: str,
    kind: SlotKind = "filter",
    field_key: FieldKey | None = None,
    field_name: str | None = None,
) -> Iterator[_Slot]:
    """Criteria of an ``<autoFilter>`` block: ``filter val=`` and ``customFilter val=``.

    Used for sheet autofilters, worksheet-table autofilters and the autofilter
    Excel nests inside a pivot label filter (which mirrors ``stringValue1``).
    """
    for column in getattr(auto_filter, "filterColumn", None) or []:
        filters = column.filters
        if filters is not None and filters.filter:
            for i in range(len(filters.filter)):
                yield _Slot(
                    f"{surface}.col{column.colId}",
                    kind,
                    _ListItem(filters.filter, i),
                    "item",
                    field_key,
                    field_name,
                )
        custom = column.customFilters
        for custom_filter in getattr(custom, "customFilter", None) or []:
            yield _Slot(
                f"{surface}.col{column.colId}.custom",
                kind,
                custom_filter,
                "val",
                field_key,
                field_name,
            )


def _str_cache_slots(source: Any, surface: str) -> Iterator[_Slot]:
    """String points of a chart data source: strRef/strCache, strLit, multiLvlStrRef."""

    def _points(data: Any) -> Iterator[_Slot]:
        for pt in getattr(data, "pt", None) or []:
            yield _Slot(surface, "filter", pt, "v")

    if source is None:
        return
    str_ref = getattr(source, "strRef", None)
    if str_ref is not None and str_ref.strCache is not None:
        yield from _points(str_ref.strCache)
    yield from _points(getattr(source, "strLit", None))
    multi = getattr(source, "multiLvlStrRef", None)
    if multi is not None and multi.multiLvlStrCache is not None:
        for level in multi.multiLvlStrCache.lvl or []:
            yield from _points(level)


def _title_slots(title: Any, surface: str) -> Iterator[_Slot]:
    """Rich-text runs (and cached cell text) of a chart or axis title."""
    tx = getattr(title, "tx", None)
    if tx is None:
        return
    rich = getattr(tx, "rich", None)
    if rich is not None:
        for paragraph in rich.p or []:
            for run in paragraph.r or []:
                yield _Slot(surface, "free", run, "t")
    yield from _str_cache_slots(tx, surface)


def _chart_slots(charts: list[Any], where: str) -> Iterator[_Slot]:
    """Titles and string caches of every chart anchored to one (chart)sheet.

    Excel always writes a ``strCache``/``numCache`` snapshot of the charted
    cells, so an Excel-authored chart keeps rendering the original names even
    after the sheet cells were pseudonymized. ``numCache`` points are numeric
    by construction (openpyxl coerces them to float) and carry no text.
    """
    for chart_index, chart in enumerate(charts, start=1):
        base = f"chart{chart_index}"
        titled = [(f"{base}.title", getattr(chart, "title", None))]
        for axis_name in ("x_axis", "y_axis", "z_axis"):
            axis = getattr(chart, axis_name, None)
            if axis is not None:
                titled.append((f"{base}.{axis_name}.title", getattr(axis, "title", None)))
        for label, title in titled:
            if title is not None:
                yield from _title_slots(title, f"{label}@{where}")
        for series_index, series in enumerate(getattr(chart, "series", None) or [], start=1):
            sbase = f"{base}.series{series_index}"
            tx = series.tx
            if tx is not None:
                yield _Slot(f"{sbase}.name@{where}", "free", tx, "v")
                yield from _str_cache_slots(tx, f"{sbase}.name@{where}")
            yield from _str_cache_slots(series.cat, f"{sbase}.cat@{where}")
            yield from _str_cache_slots(getattr(series, "xVal", None), f"{sbase}.xVal@{where}")


def _rule_slots(ws: Any, where: str) -> Iterator[_Slot]:
    """Conditional-formatting text criteria: ``text=`` and pure string-literal formulas."""
    ordinal = 0
    for fmt in ws.conditional_formatting:
        for rule in fmt.rules:
            ordinal += 1
            text = rule.text
            if isinstance(text, str):
                yield _Slot(f"cf{ordinal}.text@{where}", "filter", _RuleText(rule), "text")
            for i, formula in enumerate(rule.formula):
                if not (isinstance(formula, str) and _QUOTED_LITERAL.match(formula)):
                    continue
                if isinstance(text, str) and formula == _formula_quote(text):
                    continue  # the _RuleText slot already rewrites this echo
                yield _Slot(
                    f"cf{ordinal}.formula@{where}", "filter", _FormulaLiteral(rule, i), "literal"
                )


def _validation_slots(ws: Any, where: str) -> Iterator[_Slot]:
    """Data validations: inline list entries plus prompt/error titles and texts."""
    for ordinal, dv in enumerate(ws.data_validations.dataValidation, start=1):
        base = f"dv{ordinal}"
        if (
            dv.type == "list"
            and isinstance(dv.formula1, str)
            and _QUOTED_LITERAL.match(dv.formula1)
        ):
            for i in range(len(_formula_unquote(dv.formula1).split(","))):
                yield _Slot(f"{base}.list@{where}", "filter", _InlineListItem(dv, i), "item")
        for attr in ("promptTitle", "prompt", "errorTitle", "error"):
            yield _Slot(f"{base}.{attr}@{where}", "free", dv, attr)


def _pivot_table_slots(
    pivot: Any, table_ordinal: int, cache_ordinal: int | None, field_names: list[str]
) -> Iterator[_Slot]:
    """User-visible strings of one pivot *table* definition (not its cache).

    Item captions, page-field names/captions and label-filter values mirror the
    cached item strings, so they share the cache field's classification (and
    the mapper already knows the values once the cache slots — yielded first —
    were processed).
    """
    base = f"pivottable{table_ordinal}"

    def _key(fld: int | None) -> FieldKey | None:
        if cache_ordinal is None or fld is None or not 0 <= fld < len(field_names):
            return None
        return (cache_ordinal, fld)

    def _name(fld: int | None) -> str | None:
        if fld is None or not 0 <= fld < len(field_names):
            return None
        return field_names[fld]

    yield _Slot(f"{base}.dataCaption", "free", pivot, "dataCaption")
    for field_index, pivot_field in enumerate(pivot.pivotFields):
        for item in getattr(pivot_field, "items", None) or []:
            yield _Slot(f"{base}.item.n", "pivot", item, "n", _key(field_index), _name(field_index))
    for page_field in pivot.pageFields:
        key, name = _key(page_field.fld), _name(page_field.fld)
        yield _Slot(f"{base}.pageField.name", "pivot", page_field, "name", key, name)
        yield _Slot(f"{base}.pageField.cap", "pivot", page_field, "cap", key, name)
    for filter_index, pivot_filter in enumerate(pivot.filters, start=1):
        key, name = _key(pivot_filter.fld), _name(pivot_filter.fld)
        surface = f"{base}.filter{filter_index}"
        yield _Slot(f"{surface}.stringValue1", "pivot", pivot_filter, "stringValue1", key, name)
        yield _Slot(f"{surface}.stringValue2", "pivot", pivot_filter, "stringValue2", key, name)
        if pivot_filter.autoFilter is not None:
            # Excel mirrors stringValue1 into a nested autofilter criterion.
            yield from _autofilter_slots(pivot_filter.autoFilter, surface, "pivot", key, name)


def iter_text_slots(wb: Workbook) -> Iterator[_Slot]:
    """Yield every string slot outside the cell grid, in a deterministic order.

    A comment's ``author`` slot is always yielded directly before its ``text`` slot;
    the redact consumer relies on that to apply the author-prefix rule.
    """
    from openpyxl.packaging.custom import StringProperty

    props = wb.properties
    for name in _CORE_AUTHOR_FIELDS:
        yield _Slot(f"core.{name}", "author", props, name)
    for name in _CORE_FREE_FIELDS:
        yield _Slot(f"core.{name}", "free", props, name)

    for prop in wb.custom_doc_props.props:
        if isinstance(prop, StringProperty):
            yield _Slot(f"custom.{prop.name}", "free", prop, "value")

    for index, ws in enumerate(wb.worksheets, start=1):
        where = f"sheet{index}"
        yield from _header_footer_slots(ws, where)
        # Every comment- or hyperlink-bearing cell already exists in ws._cells (the
        # reader attaches both there); iter_rows() would materialize the whole grid.
        for (_row, _col), cell in sorted(ws._cells.items()):
            comment = cell.comment
            if comment is not None:
                yield _Slot(
                    f"comment.author@{where}!{cell.coordinate}", "author", comment, "author"
                )
                yield _Slot(f"comment.text@{where}!{cell.coordinate}", "free", comment, "text")
            link = cell.hyperlink
            if link is not None:
                at = f"@{where}!{cell.coordinate}"
                yield _Slot(f"hyperlink.tooltip{at}", "free", link, "tooltip")
                yield _Slot(f"hyperlink.display{at}", "free", link, "display")
                target = getattr(link, "target", None)
                if isinstance(target, str) and target.lower().startswith("mailto:"):
                    yield _Slot(f"hyperlink.mailto{at}", "author", _MailtoTarget(link), "address")
        yield from _autofilter_slots(ws.auto_filter, f"autofilter@{where}")
        for table_index, table in enumerate(ws.tables.values(), start=1):
            yield from _autofilter_slots(table.autoFilter, f"table{table_index}.autofilter@{where}")
        yield from _rule_slots(ws, where)
        yield from _validation_slots(ws, where)
        yield from _chart_slots(ws._charts, where)

    for index, chartsheet in enumerate(wb.chartsheets, start=1):
        yield from _header_footer_slots(chartsheet, f"chartsheet{index}")
        yield from _chart_slots(chartsheet._charts, f"chartsheet{index}")

    # Pivot caches. openpyxl exposes a sheet's pivot tables only via the private
    # ``ws._pivots`` and typed cache items only via ``._fields`` — there is no public
    # accessor. Pivot tables on ONE sheet share a cache object; across sheets the
    # reader hands every sheet its own copy of the same cacheId, and the writer emits
    # every copy. So every object gets its placeholders, but each cacheId is counted
    # and classified once.
    ordinals: dict[int, int] = {}
    seen_objects: set[int] = set()
    table_ordinal = 0
    for ws in wb.worksheets:
        for pivot in ws._pivots:
            table_ordinal += 1
            cache = pivot.cache
            cache_ordinal: int | None = None
            field_names: list[str] = []
            if cache is not None:
                field_names = [f.name for f in cache.cacheFields]
                cache_id = pivot.cacheId if pivot.cacheId is not None else id(cache)
                primary = cache_id not in ordinals
                if primary:
                    ordinals[cache_id] = len(ordinals) + 1
                cache_ordinal = ordinals[cache_id]
                if id(cache) not in seen_objects:
                    seen_objects.add(id(cache))
                    yield from _pivot_cache_slots(cache, cache_ordinal, primary)
            # Table-level captions/filters come AFTER the cache slots so the redact
            # pass has already taught the mapper the item strings they mirror.
            yield from _pivot_table_slots(pivot, table_ordinal, cache_ordinal, field_names)


def _pivot_cache_slots(cache: Any, ordinal: int, primary: bool) -> Iterator[_Slot]:
    """String slots of one pivot cache: refreshedBy, shared/group/record items."""
    from openpyxl.pivot.fields import Text as PivotText

    yield _Slot(f"pivot{ordinal}.refreshedBy", "author", cache, "refreshedBy", count=primary)
    records = list(cache.records.r) if cache.records is not None else []
    for field_index, cache_field in enumerate(cache.cacheFields):
        key = (ordinal, field_index)
        surface = f"pivot{ordinal}.{cache_field.name}"
        shared = cache_field.sharedItems
        items = list(shared._fields) if shared is not None else []
        group = cache_field.fieldGroup
        if group is not None and group.groupItems is not None:
            items.extend(group.groupItems.s)
        items.extend(r._fields[field_index] for r in records if field_index < len(r._fields))
        for item in items:
            if not isinstance(item, PivotText):
                continue
            yield _Slot(surface, "pivot", item, "v", key, cache_field.name, primary)
            if isinstance(getattr(item, "c", None), str):
                yield _Slot(surface, "pivot", item, "c", key, cache_field.name, primary)


def _author_entity_type(value: str) -> str:
    return "EMAIL" if "@" in value else "PERSON"


def _placeholder_type(placeholder: str, default: str) -> str:
    match = _PLACEHOLDER_TYPE.match(placeholder)
    return match.group(1) if match else default


async def _detect_all(
    texts: list[str], detector: DetectorLike, language: str
) -> list[list[DetectedEntity]]:
    sem = asyncio.Semaphore(8)

    async def _one(text: str) -> list[DetectedEntity]:
        async with sem:
            return await detector.detect(text, language)

    return await asyncio.gather(*[_one(t) for t in texts])


async def _classify_pivot_fields(
    slots: list[_Slot],
    detector: DetectorLike,
    language: str,
    sample_size: int,
    known_fields: Mapping[str, str],
    result: PartsResult,
) -> dict[FieldKey, str]:
    """Classify each pivot field like a sheet column.

    Tier 0: the sheet pass already classified a column of the same name.
    Tier 1: keyword on the field name. Tier 2: detector sample of the first items.
    """
    field_types: dict[FieldKey, str] = {}
    surfaces: dict[FieldKey, str] = {}
    samples: list[tuple[FieldKey, str]] = []
    per_field: dict[FieldKey, int] = {}

    for slot in slots:
        if slot.kind != "pivot" or slot.field_key is None:
            continue
        key = slot.field_key
        if key not in surfaces:
            surfaces[key] = slot.surface
            name = (slot.field_name or "").strip().lower()
            if name in known_fields:
                field_types[key] = known_fields[name]
                result.classifications[slot.surface] = f"{known_fields[name]} (sheet)"
            else:
                entity_type = infer_entity_type(slot.field_name)
                if entity_type is not None:
                    field_types[key] = entity_type
                    result.classifications[slot.surface] = f"{entity_type} (header)"
        if key in field_types or slot.attr != "v" or slot.value is None:
            continue
        if per_field.get(key, 0) < sample_size:
            samples.append((key, slot.value))
            per_field[key] = per_field.get(key, 0) + 1

    for key, entity_type in (await classify_by_sample(samples, detector, language)).items():
        field_types[key] = entity_type
        result.classifications[surfaces[key]] = f"{entity_type} (sampled)"
    return field_types


async def pseudonymize_workbook_parts(
    wb: Workbook,
    detector: DetectorLike,
    mapper: PseudonymMapper,
    language: str,
    *,
    apply: bool = True,
    sample_size: int = 5,
    known_fields: Mapping[str, str] | None = None,
) -> PartsResult:
    """Pseudonymize every part-level string slot of *wb* in place.

    *known_fields* maps lower-cased column headers the sheet pass classified to
    their entity type, so a pivot field built from that column is treated the
    same way even when its own sample would miss.

    With ``apply=False`` (detect-only / block modes) entities are detected and
    counted but nothing is written and the mapper is never touched — the proxy
    shares one mapper between message text and files.
    """
    from noirdoc.pseudonymization.engine import PseudonymizationEngine

    result = PartsResult()
    slots = [s for s in iter_text_slots(wb) if s.value is not None]

    # Detect each distinct free text once — thousands of identical "ok" comments are common.
    # "filter" slots fall back to the detector when the mapper does not know them.
    free_indexes = [i for i, s in enumerate(slots) if s.kind in ("free", "filter")]
    unique_texts: dict[str, int] = {}
    for i in free_indexes:
        text = slots[i].value
        if text is not None and text not in unique_texts:
            unique_texts[text] = len(unique_texts)
    unique_results = await _detect_all(list(unique_texts), detector, language)
    detected = {
        i: unique_results[unique_texts[text]]
        for i in free_indexes
        if (text := slots[i].value) is not None
    }
    field_types = await _classify_pivot_fields(
        slots, detector, language, sample_size, known_fields or {}, result
    )
    engine = PseudonymizationEngine()

    # comment object id -> (original author, entity type, placeholder or None when counting only)
    authors: dict[int, tuple[str, str, str | None]] = {}

    for index, slot in enumerate(slots):
        value = slot.value
        if value is None:
            continue

        if slot.kind == "author":
            if _THREADED_MIRROR_AUTHOR.match(value.strip()):
                continue  # Excel's synthetic author for threaded-comment mirrors, not PII
            if mapper.reverse_lookup(value) is not None:
                continue  # already one of this mapper's placeholders
            entity_type = _author_entity_type(value)
            placeholder = mapper.get_or_create(value, entity_type) if apply else None
            if placeholder is not None:
                slot.set(placeholder)
            authors[id(slot.obj)] = (value, entity_type, placeholder)
            if slot.count:
                result._add(entity_type)
                result.classifications[slot.surface] = f"{entity_type} (forced)"
            continue

        if slot.kind == "pivot":
            core = value.strip()
            if mapper.reverse_lookup(core) is not None:
                continue
            field_type = field_types.get(slot.field_key) if slot.field_key else None
            existing = mapper.lookup(core)  # the cell pass may know it even if the field does not
            if field_type is None and existing is None:
                continue
            item_type = field_type or _placeholder_type(existing or "", "PERSON")
            if apply:
                lead = value[: len(value) - len(value.lstrip())]
                trail = value[len(value.rstrip()) :]
                placeholder = existing or mapper.get_or_create(core, item_type)
                slot.set(lead + placeholder + trail)
            if slot.count:
                result._add(item_type)
            continue

        if slot.kind == "filter":
            # Criteria and captions mirror cell values, and the cell pass ran first:
            # a value the mapper knows gets the SAME placeholder, keeping the
            # filter/criterion consistent with the rewritten cells.
            core = value.strip()
            if mapper.reverse_lookup(core) is not None:
                continue  # already one of this mapper's placeholders
            known = mapper.lookup(core)
            if known is not None:
                if apply:
                    lead = value[: len(value) - len(value.lstrip())]
                    trail = value[len(value.rstrip()) :]
                    slot.set(lead + known + trail)
                item_type = _placeholder_type(known, "PERSON")
                result._add(item_type)
                result.classifications.setdefault(slot.surface, f"{item_type} (mapped)")
                continue
            # Unknown value: fall through to the detector path below, which keeps
            # wildcards and operators around the matched spans intact.

        if slot.kind in ("free", "filter"):
            entities = detected.get(index, [])
            # Excel's legacy comment convention: the body starts with "Author:". NER is not
            # reliable on a bare name followed by a colon — and a partial span over it would
            # leak the rest — so the prefix is replaced deterministically and the detector's
            # spans inside it are dropped BEFORE the offset-based substitution runs.
            prefix_info = authors.get(id(slot.obj)) if slot.attr == "text" else None
            prefixed = prefix_info is not None and value.startswith(f"{prefix_info[0]}:")
            if prefixed and prefix_info is not None:
                author_len = len(prefix_info[0])
                entities = [e for e in entities if e.start >= author_len]

            new_value = value
            if entities:
                if apply:
                    new_value = engine.pseudonymize(value, entities, mapper)
                for entity in entities:
                    result._add(entity.entity_type)
                best = max(entities, key=lambda e: e.score)
                result.classifications[slot.surface] = f"{best.entity_type} (detected)"

            if prefixed and prefix_info is not None:
                author, author_type, placeholder = prefix_info
                if placeholder is not None:
                    new_value = placeholder + new_value[len(author) :]
                result._add(author_type)
                result.classifications.setdefault(slot.surface, f"{author_type} (forced)")

            if apply and new_value != value:
                slot.set(new_value)

    logger.debug("xlsx_parts.completed", entity_types=result.entity_types, apply=apply)
    return result


def reidentify_workbook_parts(
    wb: Workbook, engine: ReidentificationEngine, mapper: PseudonymMapper
) -> bool:
    """Reverse placeholders in every part-level slot. Returns ``True`` if anything changed."""
    changed = False
    for slot in iter_text_slots(wb):
        value = slot.value
        if value is None or not _PLACEHOLDER.search(value):
            continue
        revealed = engine.reidentify(value, mapper)
        if revealed != value:
            slot.set(revealed)
            changed = True
    return changed


_DC_NS = "{http://purl.org/dc/elements/1.1/}"
_APP_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/extended-properties}"
_PERSONS_NS = "{http://schemas.microsoft.com/office/spreadsheetml/2018/threadedcomments}"
_APP_FIELDS = (("Manager", "PERSON"), ("Company", "ORGANIZATION"))


def clear_phantom_creator(wb: Workbook, data: bytes) -> None:
    """Undo openpyxl's default ``creator="openpyxl"`` when the file has no ``dc:creator``.

    The reader substitutes its constructor default for an absent element, which
    would otherwise be reported and mapped as a PERSON that was never in the file.
    """
    from openpyxl.xml.functions import fromstring

    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            if "docProps/core.xml" not in zf.namelist():
                wb.properties.creator = None
                return
            root = fromstring(zf.read("docProps/core.xml"))
    except Exception as exc:
        logger.warning("xlsx_parts.core_xml_unreadable", error=str(exc))
        return
    if root.find(f"{_DC_NS}creator") is None:
        wb.properties.creator = None


def count_unsupported_part_pii(data: bytes) -> PartsResult:
    """Count PII in parts openpyxl cannot model and therefore DROPS on save.

    ``docProps/app.xml`` (Manager, Company) is rebuilt blank and ``xl/persons/*``
    (threaded-comment identities) is not written at all — but only when a save
    happens at all. Counting these hits makes a workbook whose only PII lives
    there trigger the rewrite, and keeps detect/block modes honest. Nothing is
    mapped, so there is nothing to reveal.
    """
    from openpyxl.xml.functions import fromstring  # entity-safe parser openpyxl uses itself

    def _parse(zf: zipfile.ZipFile, name: str) -> Any:
        # These parts are dropped on write whatever they contain, so a corrupt one must
        # not fail a request that would otherwise redact fine.
        try:
            return fromstring(zf.read(name))
        except Exception as exc:
            logger.warning("xlsx_parts.part_unreadable", part=name, error=str(exc))
            return None

    result = PartsResult()
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        names = set(zf.namelist())
        root = _parse(zf, "docProps/app.xml") if "docProps/app.xml" in names else None
        if root is not None:
            for tag, entity_type in _APP_FIELDS:
                text = root.findtext(f"{_APP_NS}{tag}")
                if isinstance(text, str) and text.strip():
                    result._add(entity_type)
                    result.classifications[f"app.{tag}"] = f"{entity_type} (dropped)"
        person_ordinal = 0
        for name in sorted(n for n in names if n.startswith("xl/persons/") and n.endswith(".xml")):
            root = _parse(zf, name)
            if root is None:
                continue
            for person in root.iter(f"{_PERSONS_NS}person"):
                display_name = person.get("displayName")
                if isinstance(display_name, str) and display_name.strip():
                    person_ordinal += 1
                    result._add("PERSON")
                    result.classifications[f"persons.{person_ordinal}"] = "PERSON (dropped)"
    return result
